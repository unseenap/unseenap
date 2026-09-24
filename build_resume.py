from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = r"C:\Users\abhi1\Documents\Unseenap\Abhishek_Prajapati_Resume_Updated.docx"


def add_hyperlink(paragraph, text, url, bold=False):
    part = paragraph.part
    rel_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1155CC")
    props.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(underline)
    if bold:
        props.append(OxmlElement("w:b"))
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "19")
    props.append(size)
    run.append(props)
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.append(text_element)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def set_run(run, size=9, bold=False, color="000000", italic=False):
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def compact(paragraph, before=0, after=0, line=1.0):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def section_heading(doc, text):
    p = doc.add_paragraph(style="Resume Section")
    compact(p, before=5, after=2)
    run = p.add_run(text.upper())
    set_run(run, size=10.8, bold=True)
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="Resume Bullet")
    compact(p, after=0.7, line=1.02)
    p.paragraph_format.left_indent = Inches(0.16)
    p.paragraph_format.first_line_indent = Inches(-0.13)
    run = p.add_run("• ")
    set_run(run, size=9.2)
    run = p.add_run(text)
    set_run(run, size=9.2)
    return p


def entry_header(doc, title, tech=None, date=None, github=None, live=None):
    p = doc.add_paragraph()
    compact(p, before=2.1, after=0.55)
    if date:
        p.paragraph_format.tab_stops.add_tab_stop(Inches(7.1), WD_TAB_ALIGNMENT.RIGHT)
    run = p.add_run(title)
    set_run(run, size=9.7, bold=True)
    if tech:
        run = p.add_run(" | " + tech)
        set_run(run, size=9.1, italic=True, color="333333")
    if date:
        run = p.add_run("\t" + date)
        set_run(run, size=9.2, bold=True)
    if github or live:
        run = p.add_run("  ")
        set_run(run, size=9.2)
        if github:
            add_hyperlink(p, "GitHub", github)
        if live:
            run = p.add_run(" | ")
            set_run(run, size=9.2, color="555555")
            add_hyperlink(p, "Live", live)
    return p


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.5)
section.bottom_margin = Inches(0.5)
section.left_margin = Inches(0.55)
section.right_margin = Inches(0.55)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Arial"
normal.font.size = Pt(9.5)
normal.font.color.rgb = RGBColor(0, 0, 0)

if "Resume Section" not in styles:
    section_style = styles.add_style("Resume Section", WD_STYLE_TYPE.PARAGRAPH)
    section_style.font.name = "Arial"
    section_style.font.size = Pt(10.8)
    section_style.font.bold = True
    section_style.font.color.rgb = RGBColor(0, 0, 0)

if "Resume Bullet" not in styles:
    bullet_style = styles.add_style("Resume Bullet", WD_STYLE_TYPE.PARAGRAPH)
    bullet_style.font.name = "Arial"
    bullet_style.font.size = Pt(9.2)

name = doc.add_paragraph()
name.alignment = WD_ALIGN_PARAGRAPH.CENTER
compact(name, after=0.5)
run = name.add_run("ABHISHEK PRAJAPATI")
set_run(run, size=20, bold=True)

tagline = doc.add_paragraph()
tagline.alignment = WD_ALIGN_PARAGRAPH.CENTER
compact(tagline, after=0.8)
run = tagline.add_run("Full-Stack Developer | Backend Development | Workflow Automation")
set_run(run, size=10, bold=True, color="333333")

contact = doc.add_paragraph()
contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
compact(contact, after=2)
run = contact.add_run("Greater Noida, India | +91-8887878170 | ")
set_run(run, size=9.2)
add_hyperlink(contact, "Email", "mailto:abhi.prajapati2005@gmail.com")
run = contact.add_run(" | ")
set_run(run, size=9.2)
add_hyperlink(contact, "LinkedIn", "https://www.linkedin.com/in/abhishek-prajapati-9b049728a/")
run = contact.add_run(" | ")
set_run(run, size=9.2)
add_hyperlink(contact, "GitHub", "https://github.com/unseenap")

section_heading(doc, "Professional Summary")
p = doc.add_paragraph()
compact(p, after=1, line=1.02)
run = p.add_run(
    "Final-year Computer Science student and full-stack developer experienced in building production websites, "
    "role-based applications, real-time systems and workflow automation. Strong foundation in API design, "
    "database modelling, authentication, testing, deployment and SEO-focused web development."
)
set_run(run, size=9.3)

section_heading(doc, "Technical Skills")
skills = [
    ("Languages", "JavaScript, TypeScript, Python, Java, PHP, C/C++"),
    ("Frontend", "React, Next.js, Angular, HTML5, CSS3, Tailwind CSS, responsive design, SEO"),
    ("Backend and Data", "Node.js, Express.js, Spring Boot, REST APIs, Socket.IO, MySQL, PostgreSQL, MongoDB, Drizzle ORM"),
    ("Tools", "Git, GitHub Actions, Docker, Postman, AWS, Vercel, Netlify, Render, Cloudinary, Figma"),
]
for label, value in skills:
    p = doc.add_paragraph()
    compact(p, after=0.35, line=1.0)
    run = p.add_run(label + ": ")
    set_run(run, size=9.2, bold=True)
    run = p.add_run(value)
    set_run(run, size=9.2)

section_heading(doc, "Experience")
entry_header(
    doc,
    "Web Developer Intern — Cawnpore Engineering Services",
    "React.js, JavaScript, SEO",
    "Jun–Jul 2026",
    "https://github.com/unseenap/CawnporeEngineeringServices",
    "https://www.ceservices.co.in/",
)
bullet(doc, "Developed and launched the company’s production website for its HVAC design, installation, commissioning, maintenance and upgrade services across India.")
bullet(doc, "Built a responsive and accessible React interface for services, industries and completed projects; implemented on-page SEO and managed the deployed site.")

section_heading(doc, "Selected Projects")
entry_header(
    doc,
    "Auto-Examination Management System",
    "PHP, MySQL, JavaScript",
    "Mar 2026",
    "https://github.com/unseenap/Auto-Exam-Main",
)
bullet(doc, "Automated exam scheduling, multi-branch seat allocation, invigilation and attendance through Admin, Exam Cell and Faculty workflows; used validated CSV imports and transaction-safe relational operations.")

entry_header(
    doc,
    "Bodhi-Mitra",
    "React 19, Node.js, Socket.IO, MongoDB",
    "Aug 2025",
    "https://github.com/unseenap/Bodhi-Mitra",
    "https://bodhimitra.netlify.app/",
)
bullet(doc, "Built a real-time mental-health support platform with OTP-verified authentication, three-role RBAC, emergency matching, live sessions, notifications and property-based tests for critical state transitions.")

entry_header(
    doc,
    "CredX",
    "Next.js 16, TypeScript, MongoDB, Groq AI",
    "Jul 2026",
    "https://github.com/unseenap/CredX-SmartJobMatchingDash-master",
    "https://cred-x-smart-job-matching-dash.vercel.app/",
)
bullet(doc, "Created an explainable job-matching workspace with transparent 0–100 scoring, PDF/DOCX/image résumé processing, application tracking and separate student and recruiter workflows.")

entry_header(
    doc,
    "CarouselQueue",
    "Python 3.12, Instagram Graph API, Cloudinary, GitHub Actions",
    "Sep 2026",
    "https://github.com/unseenap/CarouselQueue",
)
bullet(doc, "Engineered a reusable automation pipeline that validates queued content, assembles four-slide carousels, publishes scheduled posts and records idempotent posting state; added dry-run previews and automated tests.")

entry_header(
    doc,
    "Intelligent Land Record Digitization",
    "Next.js, TypeScript, PostgreSQL, Drizzle, Zod",
    "Sep 2026",
    "https://github.com/unseenap/Land_dizitization",
)
bullet(doc, "Developed a secure, role-scoped document workflow preserving source files, OCR evidence, corrections and immutable approvals; implemented durable processing, validation, duplicate review, audit history and human verification boundaries.")

section_heading(doc, "Education")
p = doc.add_paragraph()
compact(p, after=0.25)
p.paragraph_format.tab_stops.add_tab_stop(Inches(7.1), WD_TAB_ALIGNMENT.RIGHT)
run = p.add_run("Bachelor of Technology in Computer Science and Engineering")
set_run(run, size=9.5, bold=True)
run = p.add_run("\tExpected 2027")
set_run(run, size=9.2, bold=True)
p = doc.add_paragraph()
compact(p, after=0.2)
run = p.add_run("Gautam Buddha University, Greater Noida | CGPA: 7.94/10")
set_run(run, size=9.2)

section_heading(doc, "Achievements")
p = doc.add_paragraph()
compact(p, after=0)
run = p.add_run("1st Prize — AI Innovation Hackathon 2026  |  Finalist — India Innovates 2026  |  Web Development Certification")
set_run(run, size=9.2)

core = doc.core_properties
core.title = "Abhishek Prajapati Resume"
core.subject = "Full-stack development, backend engineering and workflow automation"
core.author = "Abhishek Prajapati"

doc.save(OUTPUT)
print(OUTPUT)
