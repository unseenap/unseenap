from copy import deepcopy
import hashlib
import shutil

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


SOURCE = r"C:\Users\abhi1\Desktop\Abhishek_Prajapati_Resume_Dev.docx"
OUTPUT = r"C:\Users\abhi1\Documents\Unseenap\Abhishek_Prajapati_Resume_Revised.docx"
EXPECTED_HASH = "b5bc227c6eb21f2a61857ff8cda66cfe859f7a11ab3a8756ba3f356fb01bc3bd"


def sha256(path):
    with open(path, "rb") as stream:
        return hashlib.sha256(stream.read()).hexdigest()


if sha256(SOURCE) != EXPECTED_HASH:
    raise RuntimeError("The reference resume changed after template inspection.")

shutil.copyfile(SOURCE, OUTPUT)
doc = Document(OUTPUT)
source_paragraphs = list(doc.paragraphs)

prototypes = {
    "heading": source_paragraphs[3],
    "body": source_paragraphs[4],
    "skill": source_paragraphs[6],
    "education_header": source_paragraphs[12],
    "education_detail": source_paragraphs[13],
    "entry": source_paragraphs[17],
    "bullet": source_paragraphs[18],
    "link": source_paragraphs[22],
}


def strip_after_contact():
    body = doc._element.body
    for paragraph in source_paragraphs[3:]:
        body.remove(paragraph._p)


def copy_paragraph_properties(paragraph, prototype):
    if paragraph._p.pPr is not None:
        paragraph._p.remove(paragraph._p.pPr)
    if prototype._p.pPr is not None:
        paragraph._p.insert(0, deepcopy(prototype._p.pPr))
    paragraph.style = prototype.style


def paragraph_from(prototype_name):
    paragraph = doc.add_paragraph()
    copy_paragraph_properties(paragraph, prototypes[prototype_name])
    return paragraph


def format_run(run, *, bold=None, italic=None, size=None, color=None):
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def add_hyperlink(paragraph, label, url):
    relationship_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    properties.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.append(underline)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "20")
    properties.append(size)
    run.append(properties)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_heading(text):
    p = paragraph_from("heading")
    p.add_run(text)
    return p


def add_body(text):
    p = paragraph_from("body")
    p.add_run(text)
    return p


def add_labelled_bullet(label, value):
    p = paragraph_from("skill")
    r = p.add_run(label + ": ")
    r.bold = True
    p.add_run(value)
    return p


def add_bullet(text, final=False):
    p = paragraph_from("bullet")
    p.add_run(text)
    if final:
        p.paragraph_format.space_after = Pt(5)
    return p


def add_entry(title, technology, date):
    p = paragraph_from("entry")
    r = p.add_run(title + " | ")
    r.bold = True
    r = p.add_run(technology)
    format_run(r, italic=True, size=10, color="404040")
    if date:
        r = p.add_run("\t" + date)
        format_run(r, bold=True, italic=True, size=10, color="404040")
    return p


def add_links(github, live=None, label="GitHub"):
    p = paragraph_from("link")
    r = p.add_run(label + ": ")
    format_run(r, bold=True, size=10)
    add_hyperlink(p, github, github)
    if live:
        r = p.add_run("  |  Live: ")
        format_run(r, bold=True, size=10)
        add_hyperlink(p, live, live)
    return p


strip_after_contact()

# Update the existing role line without changing its layout.
role = doc.paragraphs[1]
for run in role.runs:
    run.text = ""
role.runs[0].text = "Full-Stack Developer | Backend Development & Workflow Automation"

add_heading("Professional Summary")
add_body(
    "Final-year Computer Science student and full-stack developer who enjoys turning manual workflows into dependable web applications. "
    "I have built and deployed role-based platforms using React, Node.js, PHP, SQL, MongoDB, and PostgreSQL, with hands-on work in authentication, real-time features, testing, automation, and SEO. "
    "I recently completed a paid web development internship where I delivered and managed a production website for an engineering company."
)

add_heading("Technical Skills")
add_labelled_bullet("Databases & SQL", "MySQL, PostgreSQL, MongoDB, relational schema design, SQL queries, transactions, Drizzle ORM")
add_labelled_bullet("Frontend Development", "JavaScript, TypeScript, React, Next.js, Angular, HTML5, CSS3, Tailwind CSS, responsive design, SEO")
add_labelled_bullet("Backend Development", "Node.js, Express.js, Spring Boot, PHP, REST APIs, Socket.IO, authentication and role-based access")
add_labelled_bullet("Programming Languages", "JavaScript, TypeScript, Python, Java, PHP, C/C++")
add_labelled_bullet("Developer Tools & Cloud", "Git, GitHub Actions, Docker, Postman, AWS, Cloudinary, Netlify, Vercel, Render, Figma, VS Code")

add_heading("Experience")
add_entry("Web Developer Intern - Cawnpore Engineering Services", "React.js · JavaScript · SEO", "Jun - Jul 2026")
add_bullet("Developed and launched the company’s website for its HVAC design, installation, commissioning, maintenance, and system-upgrade services across India.")
add_bullet("Created responsive pages for services, industries, completed projects, and the company’s engineering approach, keeping the experience clear and accessible across devices.")
add_bullet("Handled on-page SEO, deployment updates, and ongoing website management during the two-month paid internship.", final=True)
add_links(
    "https://github.com/unseenap/CawnporeEngineeringServices",
    "https://www.ceservices.co.in/",
)

add_heading("Education")
p = paragraph_from("education_header")
r = p.add_run("Gautam Buddha University")
r.bold = True
r = p.add_run("\tGreater Noida, India")
r.bold = True
p = paragraph_from("education_detail")
r = p.add_run("Bachelor of Technology - Computer Science and Engineering")
format_run(r, italic=True, size=10)
r = p.add_run("\t2023 - 2027 (Expected)")
format_run(r, italic=True, size=10)
add_bullet("CGPA: 7.94")
add_bullet("Relevant Coursework: Data Structures, Algorithms, Operating Systems, Networking, and Databases", final=True)

add_heading("Projects")
add_entry("Auto-Examination Management System", "PHP · MySQL · JavaScript · XAMPP", "Mar 2026")
add_bullet("Built a web platform that brings exam scheduling, seat allocation, invigilation, attendance, and reporting into one role-based workflow for Admin, Exam Cell, and Faculty users.")
add_bullet("Designed multi-branch seating allocation with manual adjustments, validated CSV imports, and a normalized MySQL schema with transaction-safe operations and report exports.", final=True)
add_links("https://github.com/unseenap/Auto-Exam-Main")

add_entry("Bodhi-Mitra", "React 19 · Node.js · Socket.IO · MongoDB", "Aug 2025")
add_bullet("Developed a real-time platform that connects students with psychologists through emergency matching, live sessions, notifications, and role-specific dashboards.")
add_bullet("Implemented OTP-verified authentication, JWT-based access control, crisis-keyword detection, and property-based tests for authentication, chat ordering, and emergency-state transitions.", final=True)
add_links("https://github.com/unseenap/Bodhi-Mitra", "https://bodhimitra.netlify.app/")

add_entry("CredX", "Next.js 16 · TypeScript · MongoDB · Groq AI", "Jul 2026")
add_bullet("Created a student and recruiter workspace that ranks opportunities using skills, GPA, and work-authorization signals, while showing how each factor contributes to the 0-100 match score.")
add_bullet("Added resume processing for PDF, DOCX, and image files along with application tracking and recruiter-management workflows.", final=True)
add_links(
    "https://github.com/unseenap/CredX-SmartJobMatchingDash-master",
    "https://cred-x-smart-job-matching-dash.vercel.app/",
)

add_entry("CarouselQueue", "Python 3.12 · Instagram Graph API · Cloudinary · GitHub Actions", "Sep 2026")
add_bullet("Built an automation pipeline that reads from two content queues, validates the required images, assembles four-slide carousels, uploads media, and publishes scheduled posts through the Instagram API.")
add_bullet("Added dry-run previews, automated tests, alternating queue selection, duplicate-post prevention, and local publishing-state tracking without storing credentials in the repository.", final=True)
add_links("https://github.com/unseenap/CarouselQueue")

add_entry("Intelligent Land Record Digitization", "Next.js · TypeScript · PostgreSQL · Drizzle", "Sep 2026")
add_bullet("Worked on a role-scoped document workflow that keeps source files, OCR evidence, corrections, approvals, and audit history connected throughout the land-record digitization process.")
add_bullet("Implemented validation, duplicate review, durable processing jobs, immutable approval snapshots, scoped search, and human-verification boundaries; live OCR and government integrations remain separate deployment work.", final=True)
add_links("https://github.com/unseenap/Land_dizitization")

add_heading("Achievements & Certifications")
add_bullet("1st Prize, AI Innovation Hackathon 2026")
add_bullet("Finalist, India Innovates 2026")
add_bullet("Completed Web Development Certification - Certificate No: UC-9e4a3cde-a9b9-4d28-8f74-45694ef5fceb")
add_links("https://ude.my/UC-9e4a3cde-a9b9-4d28-8f74-45694ef5fceb", label="Certificate URL")

doc.core_properties.title = "Abhishek Prajapati Resume"
doc.core_properties.subject = "Full-stack development, backend development, and workflow automation"
doc.core_properties.author = "Abhishek Prajapati"
doc.save(OUTPUT)

if sha256(SOURCE) != EXPECTED_HASH:
    raise RuntimeError("The source resume was modified during editing.")

print(OUTPUT)
